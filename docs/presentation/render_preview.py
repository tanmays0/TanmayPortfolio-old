#!/usr/bin/env python3
"""Rasterize PPTX slides to PNG for visual QA (approximate)."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

SRC = Path(__file__).resolve().parent / "TanmayPortfolio-Assignments-1-6.pptx"
OUT = Path(__file__).resolve().parent / "preview"
SCALE = 1920 / 13.333  # px per inch


def px(emu: int) -> int:
    inches = emu / 914400
    return int(inches * SCALE)


def color_of(shape):
    try:
        fill = shape.fill
        if fill.type is None:
            return None
        rgb = fill.fore_color.rgb
        return (rgb[0], rgb[1], rgb[2], 255)
    except Exception:
        return None


def load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Calibri.ttf",
        "/Library/Fonts/Arial.ttf",
        "/System/Library/Fonts/Helvetica.ttc",
    ]
    for path in candidates:
        p = Path(path)
        if p.exists():
            try:
                return ImageFont.truetype(str(p), size=max(10, size))
            except Exception:
                continue
    return ImageFont.load_default()


def render() -> None:
    prs = Presentation(str(SRC))
    W, H = px(prs.slide_width), px(prs.slide_height)
    OUT.mkdir(parents=True, exist_ok=True)
    for idx, slide in enumerate(prs.slides, 1):
        im = Image.new("RGBA", (W, H), (10, 10, 12, 255))
        draw = ImageDraw.Draw(im)
        for shape in slide.shapes:
            l, t, w, h = px(shape.left), px(shape.top), px(shape.width), px(shape.height)
            if w < 1 or h < 1:
                continue
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                blob = shape.image.blob
                pic = Image.open(BytesIO(blob)).convert("RGBA")
                pic = pic.resize((max(1, w), max(1, h)), Image.Resampling.LANCZOS)
                im.paste(pic, (l, t), pic)
                continue
            fill = color_of(shape)
            try:
                name = str(shape.auto_shape_type)
            except Exception:
                name = ""
            is_oval = "OVAL" in name.upper()
            is_round = "ROUND" in name.upper()
            if fill and shape.shape_type != MSO_SHAPE_TYPE.TEXT_BOX:
                if is_oval:
                    draw.ellipse((l, t, l + w, t + h), fill=fill)
                elif is_round:
                    rad = max(8, min(w, h) // 8)
                    draw.rounded_rectangle((l, t, l + w, t + h), radius=rad, fill=fill)
                else:
                    draw.rectangle((l, t, l + w, t + h), fill=fill)
            if shape.has_text_frame:
                y = t
                for p in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in p.runs)
                    if not text:
                        y += 16
                        continue
                    size = 14
                    bold = False
                    color = (242, 242, 244, 255)
                    if p.runs:
                        run = p.runs[0]
                        try:
                            if run.font.size:
                                size = int(run.font.size.pt * 1.15)
                        except Exception:
                            pass
                        try:
                            bold = bool(run.font.bold)
                        except Exception:
                            pass
                        try:
                            rgb = run.font.color.rgb
                            color = (rgb[0], rgb[1], rgb[2], 255)
                        except Exception:
                            pass
                    font = load_font(size, bold)
                    # crude wrap
                    max_w = max(8, w - 8)
                    words = text.split(" ")
                    line = ""
                    lines = []
                    for word in words:
                        trial = (line + " " + word).strip()
                        if font.getbbox(trial)[2] <= max_w or not line:
                            line = trial
                        else:
                            lines.append(line)
                            line = word
                    if line:
                        lines.append(line)
                    for ln in lines:
                        draw.text((l + 4, y), ln, font=font, fill=color)
                        y += size + 4
        im.convert("RGB").save(OUT / f"slide-{idx:02d}.png", quality=90)
        print("wrote", OUT / f"slide-{idx:02d}.png")


if __name__ == "__main__":
    render()
