#!/usr/bin/env python3
"""Build the Assignments 1–6 viva deck for TanmayPortfolio."""

from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
SHOTS = ROOT / "docs" / "screenshots"
ASSETS = Path(__file__).resolve().parent / "assets"
OUT = Path(__file__).resolve().parent / "TanmayPortfolio-Assignments-1-6.pptx"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# App design tokens (src/theme/colors.ts — dark theme)
BG = RGBColor(0x0A, 0x0A, 0x0C)
CARD = RGBColor(0x14, 0x14, 0x16)
SURFACE = RGBColor(0x1A, 0x1A, 0x1E)
BORDER = RGBColor(0x2A, 0x2A, 0x30)
TEXT = RGBColor(0xF2, 0xF2, 0xF4)
MUTED = RGBColor(0x8E, 0x8E, 0x98)
PRIMARY = RGBColor(0x5B, 0x8C, 0xFF)
PRIMARY_DK = RGBColor(0x3D, 0x63, 0xE8)
ON_PRIMARY = RGBColor(0x0A, 0x0A, 0x0C)
SUCCESS = RGBColor(0x3D, 0xAB, 0x8A)
DANGER = RGBColor(0xE2, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

ACCENTS = {
    1: RGBColor(0x5B, 0x8C, 0xFF),
    2: RGBColor(0x7C, 0x6C, 0xFF),
    3: RGBColor(0x3D, 0xB8, 0xC5),
    4: RGBColor(0xE8, 0x9A, 0x4A),
    5: RGBColor(0x3D, 0xAB, 0x8A),
    6: RGBColor(0xE0, 0x56, 0xA0),
}

FONT = "Calibri"
FONT_MONO = "Consolas"

TOTAL_SLIDES = 15


def rgb(color: RGBColor) -> str:
    return f"{color[0]:02X}{color[1]:02X}{color[2]:02X}"


def set_run(run, text: str, size: int, color: RGBColor, bold: bool = False, font: str = FONT) -> None:
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    run.font.italic = False


def _disable_autofit(tf) -> None:
    tf.word_wrap = True
    tf.auto_size = None


def add_text(
    slide,
    left,
    top,
    width,
    height,
    text: str,
    size: int = 14,
    color: RGBColor = TEXT,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    font: str = FONT,
    spacing: float = 1.08,
) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    _disable_autofit(tf)
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    tf.paragraphs[0].alignment = align
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = spacing
    set_run(p.add_run(), text, size, color, bold, font)


def add_paras(
    slide,
    left,
    top,
    width,
    height,
    lines: list[tuple[str, int, RGBColor, bool]],
    align=PP_ALIGN.LEFT,
    space_after: int = 6,
    font: str = FONT,
) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    _disable_autofit(tf)
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    for i, (text, size, color, bold) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.line_spacing = 1.08
        set_run(p.add_run(), text, size, color, bold, font)


def shape_fill(shape, color: RGBColor, line: RGBColor | None = None) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)


def rect(slide, l, t, w, h, fill: RGBColor, line: RGBColor | None = None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    shape_fill(s, fill, line)
    s.shadow.inherit = False
    return s


def rrect(slide, l, t, w, h, fill: RGBColor, line: RGBColor | None = None, adj: float = 0.08):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    shape_fill(s, fill, line)
    try:
        s.adjustments[0] = adj
    except Exception:
        pass
    s.shadow.inherit = False
    return s


def oval(slide, l, t, w, h, fill: RGBColor):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, l, t, w, h)
    shape_fill(s, fill)
    s.shadow.inherit = False
    return s


def set_bg(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def footer(slide, index: int, hint: str = "Tanmay Shinde  ·  React Native Portfolio") -> None:
    rect(slide, Inches(0), Inches(7.22), SLIDE_W, Inches(0.28), BG)
    add_text(slide, Inches(0.55), Inches(7.22), Inches(9.2), Inches(0.24), hint, 11, MUTED)
    add_text(
        slide,
        Inches(10.4),
        Inches(7.22),
        Inches(2.4),
        Inches(0.24),
        f"{index:02d}  /  {TOTAL_SLIDES:02d}",
        11,
        MUTED,
        align=PP_ALIGN.RIGHT,
    )


def top_bar(slide, color: RGBColor = PRIMARY) -> None:
    rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.055), color)


def kicker(slide, label: str, color: RGBColor, left=Inches(0.55), top=Inches(0.28)) -> None:
    add_text(slide, left, top, Inches(10.5), Inches(0.28), label, 12, color, True)


def heading(slide, title: str, top=Inches(0.52), width=Inches(8.2)) -> None:
    add_text(slide, Inches(0.55), top, width, Inches(0.5), title, 28, TEXT, True)


def pill(slide, l, t, w, h, text: str, fill: RGBColor, color: RGBColor, size: int = 11) -> None:
    rrect(slide, l, t, w, h, fill, adj=0.5)
    add_text(slide, l, t, w, h, text, size, color, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def evidence_row(slide, l, t, w, req: str, ev: str, accent: RGBColor) -> None:
    rrect(slide, l, t, w, Inches(0.72), CARD, BORDER, 0.12)
    rect(slide, l, t, Inches(0.07), Inches(0.72), accent)
    add_text(slide, l + Inches(0.22), t + Inches(0.08), w - Inches(0.32), Inches(0.28), req, 13, TEXT, True)
    add_text(slide, l + Inches(0.22), t + Inches(0.36), w - Inches(0.32), Inches(0.28), ev, 12, MUTED)


def round_image(src: Path, dest: Path, radius: int) -> Path:
    im = Image.open(src).convert("RGBA")
    mask = Image.new("L", im.size, 0)
    ImageDraw.Draw(mask).rounded_rectangle((0, 0, im.width - 1, im.height - 1), radius=radius, fill=255)
    im.putalpha(mask)
    dest.parent.mkdir(parents=True, exist_ok=True)
    im.save(dest)
    return dest


def device_frame(src: Path, dest: Path, width: int = 720) -> Path:
    img = Image.open(src).convert("RGBA")
    ratio = width / img.width
    img = img.resize((width, int(img.height * ratio)), Image.Resampling.LANCZOS)
    radius = int(width * 0.085)
    bezel = max(10, int(width * 0.028))
    shadow = int(width * 0.06)

    inner = Image.new("RGBA", img.size, (0, 0, 0, 0))
    inner.paste(img, (0, 0))
    mask = Image.new("L", img.size, 0)
    ImageDraw.Draw(mask).rounded_rectangle((0, 0, img.width - 1, img.height - 1), radius=radius, fill=255)
    inner.putalpha(mask)

    bezel_w, bezel_h = img.width + bezel * 2, img.height + bezel * 2
    bezel_r = radius + bezel
    chrome = Image.new("RGBA", (bezel_w, bezel_h), (0, 0, 0, 0))
    draw = ImageDraw.Draw(chrome)
    draw.rounded_rectangle((0, 0, bezel_w - 1, bezel_h - 1), radius=bezel_r, fill=(26, 26, 30, 255))
    draw.rounded_rectangle((1, 1, bezel_w - 2, bezel_h - 2), radius=bezel_r - 1, outline=(42, 42, 48, 255), width=2)
    chrome.paste(inner, (bezel, bezel), inner)

    canvas_w = bezel_w + shadow * 2
    canvas_h = bezel_h + shadow * 2
    canvas = Image.new("RGBA", (canvas_w, canvas_h), (0, 0, 0, 0))
    shade = Image.new("RGBA", (canvas_w, canvas_h), (0, 0, 0, 0))
    ImageDraw.Draw(shade).rounded_rectangle(
        (shadow + 6, shadow + 10, shadow + 6 + bezel_w, shadow + 10 + bezel_h),
        radius=bezel_r,
        fill=(0, 0, 0, 110),
    )
    shade = shade.filter(ImageFilter.GaussianBlur(14))
    canvas = Image.alpha_composite(canvas, shade)
    canvas.paste(chrome, (shadow, shadow), chrome)
    dest.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(dest, "PNG")
    return dest


def circle_avatar(src: Path, dest: Path, size: int = 220) -> Path:
    im = Image.open(src).convert("RGBA")
    im = ImageOps.fit(im, (size, size), centering=(0.5, 0.22))
    mask = Image.new("L", (size, size), 0)
    ImageDraw.Draw(mask).ellipse((0, 0, size - 1, size - 1), fill=255)
    im.putalpha(mask)
    ring = Image.new("RGBA", (size + 16, size + 16), (0, 0, 0, 0))
    ImageDraw.Draw(ring).ellipse((0, 0, size + 15, size + 15), fill=(91, 140, 255, 255))
    ring.paste(im, (8, 8), im)
    dest.parent.mkdir(parents=True, exist_ok=True)
    ring.save(dest)
    return dest


def add_phone(slide, framed: Path, left, top, height) -> None:
    slide.shapes.add_picture(str(framed), left, top, height=height)


def blank(prs) -> object:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    return slide


def new_prs() -> Presentation:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


# ---------------------------------------------------------------------------
# Slides
# ---------------------------------------------------------------------------

def slide_title(prs, frames: dict[str, Path], avatar: Path, n: int) -> None:
    s = blank(prs)
    top_bar(s)
    # glow orbs
    oval(s, Inches(-1.2), Inches(-1.4), Inches(4.2), Inches(4.2), RGBColor(0x12, 0x18, 0x2C))
    oval(s, Inches(10.6), Inches(4.6), Inches(4.0), Inches(4.0), RGBColor(0x12, 0x16, 0x22))

    pill(s, Inches(0.55), Inches(0.85), Inches(3.15), Inches(0.36), "MAD  ·  ASSIGNMENTS 1–6", SURFACE, PRIMARY, 11)
    add_text(s, Inches(0.55), Inches(1.45), Inches(8.1), Inches(1.35), "How TanmayPortfolio\ncompletes the brief.", 40, TEXT, True, spacing=1.02)
    add_text(
        s,
        Inches(0.55),
        Inches(3.0),
        Inches(7.6),
        Inches(0.7),
        "A React Native developer portfolio — built with Community CLI,\nrun on Pixel_9, mapped 1:1 to Assignments 1 through 6.",
        16,
        MUTED,
    )

    s.shapes.add_picture(str(avatar), Inches(0.55), Inches(4.05), Inches(0.62), Inches(0.62))
    add_text(s, Inches(1.32), Inches(4.08), Inches(5.5), Inches(0.28), "Tanmay Shinde", 16, TEXT, True)
    add_text(s, Inches(1.32), Inches(4.36), Inches(6.2), Inches(0.26), "B.Tech IT  ·  MIT ADT  ·  Full Stack & Software Developer", 12, MUTED)

    rrect(s, Inches(0.55), Inches(5.05), Inches(2.35), Inches(0.9), CARD, BORDER, 0.14)
    rrect(s, Inches(3.05), Inches(5.05), Inches(2.35), Inches(0.9), CARD, BORDER, 0.14)
    rrect(s, Inches(5.55), Inches(5.05), Inches(2.55), Inches(0.9), CARD, BORDER, 0.14)
    add_text(s, Inches(0.72), Inches(5.16), Inches(2.1), Inches(0.28), "STACK", 10, PRIMARY, True)
    add_text(s, Inches(0.72), Inches(5.42), Inches(2.1), Inches(0.38), "RN 0.87 · TypeScript", 13, TEXT, True)
    add_text(s, Inches(3.22), Inches(5.16), Inches(2.1), Inches(0.28), "RUN", 10, PRIMARY, True)
    add_text(s, Inches(3.22), Inches(5.42), Inches(2.1), Inches(0.38), "Pixel_9 emulator", 13, TEXT, True)
    add_text(s, Inches(5.72), Inches(5.16), Inches(2.2), Inches(0.28), "TOOLING", 10, PRIMARY, True)
    add_text(s, Inches(5.72), Inches(5.42), Inches(2.2), Inches(0.38), "CLI + Android Studio", 13, TEXT, True)

    add_phone(s, frames["home_dark"], Inches(9.55), Inches(0.45), Inches(6.55))
    footer(s, n)
    notes(
        s,
        "Open with: this is one React Native portfolio that satisfies Assignments 1–6. "
        "Community CLI project, android/ in Android Studio, Pixel_9 emulator. "
        "Then walk the six assignments with live app + these screenshots.",
    )


def slide_agenda(prs, n: int) -> None:
    s = blank(prs)
    top_bar(s)
    kicker(s, "ROADMAP", PRIMARY)
    heading(s, "Six assignments. One running app.")
    add_text(s, Inches(0.55), Inches(1.05), Inches(12), Inches(0.32), "Each card is a viva checkpoint — requirement on the left, where to point in the project on the right.", 14, MUTED)

    items = [
        (1, "Setup + Hello World", "Community CLI · android/ · Pixel_9 · first emulator run"),
        (2, "Comparison report", "Native vs Hybrid vs React Native — and why RN here"),
        (3, "Components, Props, useState", "src/components · ProjectCard props · theme + form state"),
        (4, "Folder structure + state", "Documented src/ tree · Switch + validation updates"),
        (5, "Profile app + core RN", "View, Text, Image, TextInput, Button, Switch, ScrollView"),
        (6, "Flexbox, TextInput, onPress", "Skill wrap · contact inputs · Submit handler"),
    ]
    for i, (num, title, blurb) in enumerate(items):
        col, row = i % 3, i // 3
        l = Inches(0.55 + col * 4.15)
        t = Inches(1.55 + row * 2.55)
        color = ACCENTS[num]
        rrect(s, l, t, Inches(3.95), Inches(2.32), CARD, BORDER, 0.1)
        rect(s, l, t, Inches(0.08), Inches(2.32), color)
        oval(s, l + Inches(0.28), t + Inches(0.28), Inches(0.48), Inches(0.48), color)
        add_text(s, l + Inches(0.28), t + Inches(0.34), Inches(0.48), Inches(0.36), f"{num:02d}", 13, ON_PRIMARY, True, PP_ALIGN.CENTER)
        add_text(s, l + Inches(0.28), t + Inches(0.95), Inches(3.4), Inches(0.5), title, 18, TEXT, True)
        add_text(s, l + Inches(0.28), t + Inches(1.45), Inches(3.4), Inches(0.6), blurb, 13, MUTED)
    footer(s, n)
    notes(s, "Spend 15 seconds here: six assignments, one codebase. Then dive into A1 with the emulator story.")


def slide_a1(prs, frames: dict[str, Path], n: int) -> None:
    s = blank(prs)
    top_bar(s, ACCENTS[1])
    kicker(s, "ASSIGNMENT 01  ·  Community CLI + emulator", ACCENTS[1])
    heading(s, "Setup + Hello World")
    add_text(s, Inches(0.55), Inches(1.05), Inches(7.4), Inches(0.32), "Evidence the project is a real Community CLI app that runs on the emulator.", 14, MUTED)

    rows = [
        ("Create RN project with Community CLI", "Root TanmayPortfolio · package.json → react-native 0.87.1"),
        ("android/ openable in Android Studio", "Open the android/ Gradle project — not Expo Go"),
        ("Run on emulator", "npm run android  /  Studio Run on Pixel_9"),
        ("Hello World → running portfolio", "Same CLI app; Home on Pixel_9 is the living first screen"),
    ]
    for i, (req, ev) in enumerate(rows):
        evidence_row(s, Inches(0.55), Inches(1.5 + i * 0.86), Inches(7.55), req, ev, ACCENTS[1])

    pill(s, Inches(0.55), Inches(5.05), Inches(7.55), Inches(0.42), "Demo tip   ·   CLI init  →  Metro  →  first successful emulator install", SURFACE, PRIMARY, 13)
    rrect(s, Inches(0.55), Inches(5.6), Inches(7.55), Inches(1.28), CARD, BORDER, 0.1)
    add_text(s, Inches(0.75), Inches(5.75), Inches(7.2), Inches(0.28), "What to say", 11, ACCENTS[1], True)
    add_text(
        s,
        Inches(0.75),
        Inches(6.05),
        Inches(7.15),
        Inches(0.7),
        "I used the React Native Community CLI so I get a real android/ folder. Metro bundles JS; Android Studio installs the APK on Pixel_9. The first milestone was a Hello World screen — the same app now hosts the full portfolio.",
        13,
        MUTED,
    )

    add_phone(s, frames["home_dark"], Inches(8.55), Inches(0.55), Inches(6.4))
    add_text(s, Inches(8.55), Inches(6.95), Inches(4.2), Inches(0.22), "Pixel_9  ·  running portfolio Home", 11, MUTED, align=PP_ALIGN.CENTER)
    footer(s, n)
    notes(
        s,
        "Show package.json react-native, then android/ in Studio, then this screenshot. "
        "Mention Metro Fast Refresh. Do not spend more than 45 seconds on setup.",
    )


def slide_a2_compare(prs, n: int) -> None:
    s = blank(prs)
    top_bar(s, ACCENTS[2])
    kicker(s, "ASSIGNMENT 02  ·  comparison report", ACCENTS[2])
    heading(s, "Native vs Hybrid vs React Native")
    add_text(s, Inches(0.55), Inches(1.05), Inches(12), Inches(0.3), "Same product goal. Three ways to build. RN sits in the middle: shared JS, real native views.", 14, MUTED)

    cols = [
        (
            "Native",
            ACCENTS[4],
            False,
            [
                "Kotlin / Swift + platform SDKs",
                "Best performance & OS APIs",
                "Two codebases to maintain",
                "Banking, games, camera-heavy",
            ],
        ),
        (
            "Hybrid (WebView)",
            MUTED,
            False,
            [
                "HTML / CSS / JS in a WebView",
                "Fastest reuse of a website",
                "Weaker gestures & native feel",
                "News / brochure / simple forms",
            ],
        ),
        (
            "React Native",
            PRIMARY,
            True,
            [
                "TypeScript + React → native views",
                "One UI codebase, Metro Fast Refresh",
                "Need some Gradle / native tooling",
                "This portfolio · product UIs · SaaS",
            ],
        ),
    ]
    for i, (title, accent, featured, bullets) in enumerate(cols):
        l = Inches(0.55 + i * 4.15)
        fill = SURFACE if featured else CARD
        line = accent if featured else BORDER
        rrect(s, l, Inches(1.5), Inches(3.95), Inches(5.35), fill, line, 0.08)
        if featured:
            pill(s, l + Inches(0.28), Inches(1.7), Inches(1.55), Inches(0.32), "CHOSEN", accent, ON_PRIMARY, 10)
        add_text(s, l + Inches(0.28), Inches(2.15 if featured else 1.78), Inches(3.4), Inches(0.42), title, 22, TEXT, True)
        add_text(s, l + Inches(0.28), Inches(2.6 if featured else 2.25), Inches(3.4), Inches(0.28), "Approach", 11, accent, True)
        for j, b in enumerate(bullets):
            yt = Inches((3.05 if featured else 2.7) + j * 0.85)
            rrect(s, l + Inches(0.28), yt, Inches(3.4), Inches(0.72), CARD if featured else SURFACE, None, 0.16)
            add_text(s, l + Inches(0.42), yt + Inches(0.18), Inches(3.1), Inches(0.42), b, 13, TEXT)
    footer(s, n)
    notes(
        s,
        "Native = best performance, two codebases. Hybrid = WebView, web-like. "
        "RN = React components that map to View/Text native widgets — that is why it fits this course.",
    )


def slide_a2_why(prs, n: int) -> None:
    s = blank(prs)
    top_bar(s, ACCENTS[2])
    kicker(s, "ASSIGNMENT 02  ·  section 7 — why this stack", ACCENTS[2])
    heading(s, "Why React Native for this project")

    reasons = [
        ("01", "Curriculum fit", "Functional components, Props, useState, core RN widgets, Flexbox — the assignment list is RN’s native language."),
        ("02", "Studio workflow", "Community CLI ships android/. Open it in Android Studio, run Pixel_9, show Logcat — same as a native lab."),
        ("03", "One codebase", "A full student profile (Home → Contact) without maintaining separate Kotlin and Swift apps."),
        ("04", "Industry signal", "RN is widely used for cross-platform product UI — useful in internships, not only in the viva."),
    ]
    for i, (num, title, body) in enumerate(reasons):
        col, row = i % 2, i // 2
        l = Inches(0.55 + col * 6.3)
        t = Inches(1.45 + row * 2.35)
        rrect(s, l, t, Inches(6.05), Inches(2.15), CARD, BORDER, 0.1)
        oval(s, l + Inches(0.28), t + Inches(0.32), Inches(0.55), Inches(0.55), ACCENTS[2])
        add_text(s, l + Inches(0.28), t + Inches(0.4), Inches(0.55), Inches(0.4), num, 12, WHITE, True, PP_ALIGN.CENTER)
        add_text(s, l + Inches(1.0), t + Inches(0.38), Inches(4.7), Inches(0.4), title, 20, TEXT, True)
        add_text(s, l + Inches(0.32), t + Inches(1.05), Inches(5.4), Inches(0.85), body, 14, MUTED)

    add_text(
        s,
        Inches(0.55),
        Inches(6.3),
        Inches(12.2),
        Inches(0.55),
        "Submit: export docs/assignment-2-comparison-report.md to PDF if the college requires PDF. Native wins raw performance; Hybrid wins web reuse; RN wins this course.",
        13,
        MUTED,
    )
    footer(s, n)
    notes(s, "If asked ‘why not native?’: would not teach Props/useState as directly. ‘Why not hybrid?’: WebView is not View/Text — the rubric wants RN core components.")


def slide_a3_props(prs, frames: dict[str, Path], n: int) -> None:
    s = blank(prs)
    top_bar(s, ACCENTS[3])
    kicker(s, "ASSIGNMENT 03  ·  reusable UI", ACCENTS[3])
    heading(s, "Reusable components + Props")
    add_text(s, Inches(0.55), Inches(1.05), Inches(7.6), Inches(0.3), "Data lives in src/data. Components stay dumb and reusable.", 14, MUTED)

    comps = [
        ("ProjectCard", "title, description, tags via Props"),
        ("SkillChip", "label + accent — used in Skills & cards"),
        ("ThemeToggle", "controlled Switch, no hardcoded theme"),
        ("ContactForm", "stateful form composed of ContactInput"),
    ]
    for i, (name, desc) in enumerate(comps):
        t = Inches(1.48 + i * 0.7)
        rrect(s, Inches(0.55), t, Inches(7.55), Inches(0.62), CARD, BORDER, 0.18)
        add_text(s, Inches(0.78), t + Inches(0.08), Inches(3.2), Inches(0.22), name, 12, ACCENTS[3], True, font=FONT_MONO)
        add_text(s, Inches(0.78), t + Inches(0.3), Inches(7.1), Inches(0.24), desc, 13, TEXT)

    rrect(s, Inches(0.55), Inches(4.35), Inches(7.55), Inches(2.4), SURFACE, BORDER, 0.1)
    add_text(s, Inches(0.78), Inches(4.5), Inches(7.1), Inches(0.26), "ProjectsScreen.tsx  →  data flow", 12, ACCENTS[3], True)
    add_text(
        s,
        Inches(0.78),
        Inches(4.88),
        Inches(7.1),
        Inches(1.6),
        "src/data/projects.ts\n        ↓  map()\n<ProjectCard title={...} tags={...} />\n        ↓  Details Pressable\nexpanded state lives inside the card (useState)",
        15,
        TEXT,
        font=FONT_MONO,
        spacing=1.25,
    )

    add_phone(s, frames["projects"], Inches(8.55), Inches(0.55), Inches(6.4))
    footer(s, n)
    notes(s, "Viva line 1: Data lives in src/data; components stay reusable via Props. Open Projects, tap Details to show the card expanding from internal useState.")


def slide_a3_state(prs, frames: dict[str, Path], n: int) -> None:
    s = blank(prs)
    top_bar(s, ACCENTS[3])
    kicker(s, "ASSIGNMENT 03  ·  useState", ACCENTS[3])
    heading(s, "Dynamic UI from state — not from rewrites")

    rrect(s, Inches(0.55), Inches(1.4), Inches(5.85), Inches(5.35), CARD, BORDER, 0.08)
    add_phone(s, frames["theme"], Inches(1.55), Inches(1.55), Inches(5.05))
    add_text(s, Inches(0.7), Inches(6.38), Inches(5.55), Inches(0.28), "Theme  ·  isDark in App.tsx  ·  Switch", 12, MUTED, align=PP_ALIGN.CENTER)

    rrect(s, Inches(6.6), Inches(1.4), Inches(6.15), Inches(5.35), CARD, BORDER, 0.08)
    add_phone(s, frames["form_errors"], Inches(8.05), Inches(1.55), Inches(5.05))
    add_text(s, Inches(6.75), Inches(6.38), Inches(5.85), Inches(0.28), "Form  ·  ContactForm errors  ·  useState", 12, MUTED, align=PP_ALIGN.CENTER)

    footer(s, n)
    notes(
        s,
        "Viva line 2: isDark in App.tsx updates the whole theme when Switch toggles. "
        "Then: empty Submit shows three errors — same pattern, local form state. "
        "ProjectCard Details is a third useState (expanded).",
    )


def slide_a4(prs, n: int) -> None:
    s = blank(prs)
    top_bar(s, ACCENTS[4])
    kicker(s, "ASSIGNMENT 04  ·  documented src/ tree", ACCENTS[4])
    heading(s, "Folder structure that matches the concepts")

    tree = [
        ("android/", "Native Gradle project — open in Android Studio"),
        ("App.tsx", "Root useState(isDark) + ThemeProvider + Navigation"),
        ("src/components/", "Reusable UI: ProjectCard, SkillChip, ThemeToggle…"),
        ("src/screens/", "One screen per portfolio section"),
        ("src/data/", "Static content — keep components prop-driven"),
        ("src/theme/", "colors, tokens, ThemeContext"),
        ("src/navigation/", "Bottom tabs + More stack"),
        ("docs/", "Reports + screenshots for submission"),
    ]
    for i, (path, why) in enumerate(tree):
        col, row = i % 2, i // 2
        l = Inches(0.55 + col * 6.3)
        t = Inches(1.4 + row * 1.18)
        rrect(s, l, t, Inches(6.05), Inches(1.05), CARD, BORDER, 0.14)
        rect(s, l, t, Inches(0.08), Inches(1.05), ACCENTS[4])
        add_text(s, l + Inches(0.28), t + Inches(0.16), Inches(5.5), Inches(0.3), path, 16, TEXT, True, font=FONT_MONO)
        add_text(s, l + Inches(0.28), t + Inches(0.5), Inches(5.5), Inches(0.38), why, 13, MUTED)

    footer(s, n)
    notes(
        s,
        "A4 is A3 plus documentation. Point at the markdown tree, then: reusable components are the same; "
        "state-based updates = Theme Switch + contact validation errors. Structure is the evidence that this is not one giant App.tsx.",
    )


def slide_a5_components(prs, frames: dict[str, Path], n: int) -> None:
    s = blank(prs)
    top_bar(s, ACCENTS[5])
    kicker(s, "ASSIGNMENT 05  ·  core RN widgets", ACCENTS[5])
    heading(s, "Student profile app — eight core components")
    add_text(s, Inches(0.55), Inches(1.05), Inches(7.6), Inches(0.3), "Portfolio sections = a profile-style app: Home, About, Skills, Projects, Experience, Certs, Resume, Contact.", 13, MUTED)

    cores = [
        ("View", "Almost every layout"),
        ("Text", "Titles, bio, errors"),
        ("Image", "ProfileCard + profile.jpg"),
        ("TextInput", "ContactInput fields"),
        ("Button", "ContactForm submit"),
        ("Switch", "ThemeToggle"),
        ("ScrollView", "ScreenContainer"),
        ("StyleSheet", "Every styled file"),
    ]
    for i, (name, where) in enumerate(cores):
        col, row = i % 2, i // 2
        l = Inches(0.55 + col * 3.85)
        t = Inches(1.48 + row * 1.18)
        rrect(s, l, t, Inches(3.7), Inches(1.05), CARD, BORDER, 0.14)
        add_text(s, l + Inches(0.22), t + Inches(0.18), Inches(3.25), Inches(0.3), name, 18, ACCENTS[5], True, font=FONT_MONO)
        add_text(s, l + Inches(0.22), t + Inches(0.52), Inches(3.25), Inches(0.36), where, 13, MUTED)

    add_phone(s, frames["home_profile"], Inches(8.55), Inches(0.55), Inches(6.4))
    footer(s, n)
    notes(s, "Walk the Home screenshot: Image (photo), Text (name/bio), View (cards), Switch at bottom, ScrollView for the page, StyleSheet for spacing. Then say the other widgets live on Contact.")


def slide_a5_gallery(prs, frames: dict[str, Path], n: int) -> None:
    s = blank(prs)
    top_bar(s, ACCENTS[5])
    kicker(s, "ASSIGNMENT 05  ·  full profile surfaces", ACCENTS[5])
    heading(s, "Every section is a profile screen")

    gallery = [
        (frames["skills"], "Skills", "SkillChip + ScrollView"),
        (frames["experience"], "Experience", "AdMax + CodSoft rail"),
        (frames["more"], "More / About", "Profile-style cards"),
        (frames["contact"], "Contact", "TextInput + Button"),
    ]
    for i, (img, title, cap) in enumerate(gallery):
        l = Inches(0.4 + i * 3.25)
        add_phone(s, img, l, Inches(1.25), Inches(5.15))
        add_text(s, l, Inches(6.45), Inches(3.05), Inches(0.26), title, 14, TEXT, True, PP_ALIGN.CENTER)
        add_text(s, l, Inches(6.72), Inches(3.05), Inches(0.24), cap, 12, MUTED, align=PP_ALIGN.CENTER)
    footer(s, n)
    notes(s, "Quick pan: Skills chips, Experience timeline from the resume, More menu, Contact form. This is the ‘student profile app’ the assignment asks for — not a toy counter.")


def slide_a6_flex(prs, frames: dict[str, Path], n: int) -> None:
    s = blank(prs)
    top_bar(s, ACCENTS[6])
    kicker(s, "ASSIGNMENT 06  ·  Flexbox", ACCENTS[6])
    heading(s, "Layouts that wrap, align, and rail")

    points = [
        ("Home CTA wrap", "Explore buttons use flexWrap so CTAs reflow instead of overflowing."),
        ("SkillCategory chips", "SkillChip row with flexWrap — tags sit on multiple lines."),
        ("Experience row + rail", "Horizontal row: timeline rail on the left, card content on the right."),
        ("Project card actions", "Details + GitHub sit in a Flexbox row inside ProjectCard."),
    ]
    for i, (title, body) in enumerate(points):
        t = Inches(1.4 + i * 1.18)
        rrect(s, Inches(0.55), t, Inches(7.45), Inches(1.05), CARD, BORDER, 0.14)
        rect(s, Inches(0.55), t, Inches(0.08), Inches(1.05), ACCENTS[6])
        add_text(s, Inches(0.85), t + Inches(0.16), Inches(6.9), Inches(0.3), title, 16, TEXT, True)
        add_text(s, Inches(0.85), t + Inches(0.5), Inches(6.9), Inches(0.4), body, 13, MUTED)

    add_phone(s, frames["skills"], Inches(8.45), Inches(0.55), Inches(6.4))
    footer(s, n)
    notes(s, "On Skills, point at wrapping chips. On Experience, point at the vertical rail + row. Flexbox is not a theory slide — it is visible in those two screens.")


def slide_a6_input(prs, frames: dict[str, Path], n: int) -> None:
    s = blank(prs)
    top_bar(s, ACCENTS[6])
    kicker(s, "ASSIGNMENT 06  ·  TextInput + onPress", ACCENTS[6])
    heading(s, "Controlled inputs, then a press that means something")

    add_phone(s, frames["contact_empty"], Inches(0.45), Inches(1.25), Inches(5.35))
    add_phone(s, frames["form_errors"], Inches(3.55), Inches(1.25), Inches(5.35))

    add_text(s, Inches(0.45), Inches(6.62), Inches(3.0), Inches(0.24), "Empty form  ·  TextInput", 12, MUTED, align=PP_ALIGN.CENTER)
    add_text(s, Inches(3.55), Inches(6.62), Inches(3.0), Inches(0.24), "Submit  ·  onPress errors", 12, MUTED, align=PP_ALIGN.CENTER)

    rrect(s, Inches(6.85), Inches(1.4), Inches(5.95), Inches(5.35), CARD, BORDER, 0.08)
    add_text(s, Inches(7.1), Inches(1.62), Inches(5.5), Inches(0.32), "What the code is doing", 14, ACCENTS[6], True)
    lines = [
        ("Controlled TextInput", "Name, email, message values live in useState. The field always shows state — never an uncontrolled buffer."),
        ("Button onPress", "ContactForm <Button onPress={handleSubmit} />. Empty submit → error state. Valid submit → Alert."),
        ("Pressable CTAs", "Home Explore buttons and ProjectCard GitHub/Details use Pressable for navigation and expand."),
        ("Validation UX", "Red borders + summary box prove state changed. That is the assignment: input handling + a press handler."),
    ]
    for i, (t, b) in enumerate(lines):
        yt = Inches(2.1 + i * 1.1)
        add_text(s, Inches(7.1), yt, Inches(5.5), Inches(0.28), t, 15, TEXT, True)
        add_text(s, Inches(7.1), yt + Inches(0.32), Inches(5.5), Inches(0.7), b, 13, MUTED)
    footer(s, n)
    notes(
        s,
        "Live demo: More → Contact. Tap Submit empty — three errors. Fill fields, Submit — Alert. "
        "Say the words TextInput, onPress, and Flexbox out loud so the examiner can tick the rubric.",
    )


def slide_rubric(prs, n: int) -> None:
    s = blank(prs)
    top_bar(s, PRIMARY)
    kicker(s, "RUBRIC ALIGNMENT", PRIMARY)
    heading(s, "Aimed at Excellent — and where to prove it")

    rows = [
        ("Functionality & Output", "All sections run. Theme Switch. Form validation. External links. Pixel_9 install."),
        ("Code Structure & Syntax", "src/ modular layout. TypeScript. Consistent component names. No one-file app."),
        ("Concept Implementation", "Props, useState, eight core components, Flexbox, TextInput, onPress — all visible."),
        ("Readability & Documentation", "File comments, docs/ reports, README, screenshot folders per assignment."),
    ]
    for i, (title, body) in enumerate(rows):
        t = Inches(1.4 + i * 1.25)
        rrect(s, Inches(0.55), t, Inches(12.2), Inches(1.12), CARD, BORDER, 0.1)
        oval(s, Inches(0.78), t + Inches(0.3), Inches(0.52), Inches(0.52), SUCCESS)
        add_text(s, Inches(0.78), t + Inches(0.38), Inches(0.52), Inches(0.36), "✓", 16, WHITE, True, PP_ALIGN.CENTER)
        add_text(s, Inches(1.5), t + Inches(0.18), Inches(10.8), Inches(0.32), title, 18, TEXT, True)
        add_text(s, Inches(1.5), t + Inches(0.54), Inches(10.8), Inches(0.42), body, 14, MUTED)
    footer(s, n)
    notes(s, "Do not read this slide. Use it if they ask ‘how does this meet Excellent?’ Point at Functionality first, then Concept Implementation.")


def slide_viva(prs, n: int) -> None:
    s = blank(prs)
    top_bar(s, PRIMARY)
    kicker(s, "SUGGESTED WALKTHROUGH  ·  3–4 MINUTES", PRIMARY)
    heading(s, "Show the app. Name the concept.")

    steps = [
        ("1", "Home", "Identity, photo, CTAs — stay on light theme first."),
        ("2", "Dark mode", "Toggle Switch at the bottom of Home. Say useState / isDark."),
        ("3", "Projects", "Props on ProjectCard. Expand Details."),
        ("4", "Skills", "Flexbox chips wrapping. SkillChip reuse."),
        ("5", "Experience", "AdMax + CodSoft from the resume. Rail + row."),
        ("6", "Contact", "Empty Submit → errors. Valid Submit → Alert."),
        ("7", "Studio", "Open android/. Mention Logcat + Metro."),
    ]
    for i, (num, title, body) in enumerate(steps):
        t = Inches(1.32 + i * 0.72)
        oval(s, Inches(0.55), t + Inches(0.08), Inches(0.48), Inches(0.48), PRIMARY)
        add_text(s, Inches(0.55), t + Inches(0.14), Inches(0.48), Inches(0.36), num, 14, ON_PRIMARY, True, PP_ALIGN.CENTER)
        add_text(s, Inches(1.22), t + Inches(0.08), Inches(2.4), Inches(0.48), title, 18, TEXT, True, anchor=MSO_ANCHOR.MIDDLE)
        add_text(s, Inches(3.7), t + Inches(0.08), Inches(8.9), Inches(0.48), body, 16, MUTED, anchor=MSO_ANCHOR.MIDDLE)
    footer(s, n)
    notes(s, "This is the live demo script. Do not open the IDE until step 7. Keep the phone in your hand / emulator in front. Two viva lines: data in src/data; isDark updates the whole theme.")


def slide_close(prs, avatar: Path, n: int) -> None:
    s = blank(prs)
    top_bar(s)
    oval(s, Inches(10.2), Inches(-1.2), Inches(4.5), Inches(4.5), RGBColor(0x12, 0x18, 0x2C))
    add_text(s, Inches(0.7), Inches(1.7), Inches(12), Inches(0.4), "QUESTIONS", 14, PRIMARY, True)
    add_text(s, Inches(0.7), Inches(2.15), Inches(12), Inches(1.1), "Happy to walk any screen\nor any file.", 40, TEXT, True)

    rrect(s, Inches(0.7), Inches(3.7), Inches(11.9), Inches(2.35), CARD, BORDER, 0.08)
    add_text(s, Inches(1.0), Inches(3.92), Inches(11.3), Inches(0.3), "Keep these two lines ready", 12, PRIMARY, True)
    add_text(
        s,
        Inches(1.0),
        Inches(4.35),
        Inches(11.3),
        Inches(0.5),
        "1.  Data lives in src/data; components stay reusable via Props.",
        18,
        TEXT,
        True,
    )
    add_text(
        s,
        Inches(1.0),
        Inches(4.95),
        Inches(11.3),
        Inches(0.7),
        "2.  isDark in App.tsx updates the whole theme when the Switch toggles.",
        18,
        TEXT,
        True,
    )

    s.shapes.add_picture(str(avatar), Inches(0.7), Inches(6.25), Inches(0.5), Inches(0.5))
    add_text(s, Inches(1.38), Inches(6.28), Inches(8), Inches(0.22), "Tanmay Shinde  ·  TanmayPortfolio", 14, TEXT, True)
    add_text(s, Inches(1.38), Inches(6.52), Inches(8), Inches(0.22), "React Native Community CLI  ·  Pixel_9  ·  Assignments 1–6", 12, MUTED)
    footer(s, n, "Thank you")
    notes(s, "Stop talking. Wait for questions. If silence, offer to open ContactForm.tsx or App.tsx.")


def prepare_assets() -> tuple[dict[str, Path], Path]:
    ASSETS.mkdir(parents=True, exist_ok=True)
    mapping = {
        # Note: some files in docs/screenshots/ are misnamed; pick by actual screen.
        "home_dark": SHOTS / "assignment-3" / "03-dark-mode.png",
        "theme": SHOTS / "assignment-3" / "02-home-theme-toggle-bottom.png",
        "projects": SHOTS / "assignment-3" / "01-projects-props-cards.png",
        "form_errors": SHOTS / "assignment-3" / "04-form-state-errors.png",
        "home_profile": SHOTS / "assignment-3" / "02-home-theme-toggle-bottom.png",
        "skills": SHOTS / "assignment-5" / "02-skills-scroll.png",
        "experience": SHOTS / "assignment-5" / "03-experience.png",
        "more": SHOTS / "assignment-5" / "04-more-menu.png",
        "contact": SHOTS / "assignment-5" / "05-contact-form.png",
        "contact_empty": SHOTS / "assignment-5" / "05-contact-form.png",
    }
    frames = {key: device_frame(src, ASSETS / f"{key}.png") for key, src in mapping.items()}
    avatar = circle_avatar(ROOT / "assets" / "images" / "profile.jpg", ASSETS / "avatar.png")
    return frames, avatar


def build() -> Path:
    frames, avatar = prepare_assets()
    prs = new_prs()
    builders = [
        lambda i: slide_title(prs, frames, avatar, i),
        lambda i: slide_agenda(prs, i),
        lambda i: slide_a1(prs, frames, i),
        lambda i: slide_a2_compare(prs, i),
        lambda i: slide_a2_why(prs, i),
        lambda i: slide_a3_props(prs, frames, i),
        lambda i: slide_a3_state(prs, frames, i),
        lambda i: slide_a4(prs, i),
        lambda i: slide_a5_components(prs, frames, i),
        lambda i: slide_a5_gallery(prs, frames, i),
        lambda i: slide_a6_flex(prs, frames, i),
        lambda i: slide_a6_input(prs, frames, i),
        lambda i: slide_rubric(prs, i),
        lambda i: slide_viva(prs, i),
        lambda i: slide_close(prs, avatar, i),
    ]
    if len(builders) != TOTAL_SLIDES:
        raise RuntimeError(f"Slide count mismatch: {len(builders)} vs {TOTAL_SLIDES}")
    for idx, fn in enumerate(builders, start=1):
        fn(idx)
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(path)
